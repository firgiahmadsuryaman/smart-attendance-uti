from datetime import datetime
import json
import os
from io import BytesIO
from flask import Blueprint, flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import login_required, current_user
from sqlalchemy.exc import IntegrityError
from sqlalchemy import or_
from ..extensions import db
from ..models import (User, StudyProgram, Semester, ClassOffering, ClassSchedule, CourseEnrollment, FaceProfile,
                      AttendanceSession, Attendance, Notification)
from ..utils import role_required, log_action, get_active_semester
from ..face_engine import (extract_embedding, compare_embedding, FaceEngineError,
                           verify_blink_liveness, verify_enrollment_liveness)

student_bp = Blueprint("student", __name__, url_prefix="/student")


def _active_enrollments():
    active = get_active_semester()
    if not active:
        return []
    return CourseEnrollment.query.join(ClassOffering).filter(
        CourseEnrollment.student_id == current_user.id,
        CourseEnrollment.status == "APPROVED",
        ClassOffering.semester_id == active.id,
    ).all()


@student_bp.route("/")
@login_required
@role_required("student")
def dashboard():
    enrollments = _active_enrollments()
    offering_ids = [e.offering_id for e in enrollments]
    sessions = AttendanceSession.query.filter(AttendanceSession.offering_id.in_(offering_ids)).order_by(AttendanceSession.opened_at.desc()).all() if offering_ids else []
    active_sessions = [s for s in sessions if s.status == "ACTIVE" and s.opened_at <= datetime.now() <= s.closes_at]
    # Only closed sessions contribute to the finalized attendance percentage.
    finalized_sessions = [s for s in sessions if s.status == "CLOSED"]
    finalized_ids = [s.id for s in finalized_sessions]
    attendance_records = Attendance.query.filter(
        Attendance.student_id == current_user.id,
        Attendance.session_id.in_(finalized_ids),
    ).all() if finalized_ids else []
    present_count = sum(1 for a in attendance_records if a.status == "PRESENT")
    late_count = sum(1 for a in attendance_records if a.status == "LATE")
    absent_count = sum(1 for a in attendance_records if a.status == "ABSENT")
    finalized_total = len(finalized_sessions)
    attendance_percent = round((present_count + late_count) * 100 / finalized_total) if finalized_total else 0
    today_schedule = []
    for enrollment in enrollments:
        for schedule in enrollment.offering.schedules:
            if schedule.weekday == datetime.now().weekday():
                today_schedule.append((enrollment.offering, schedule))
    today_schedule.sort(key=lambda row: row[1].start_time)
    return render_template("student/dashboard.html", page_title="Dashboard Mahasiswa",
        enrollments=enrollments, active_sessions=active_sessions, today_schedule=today_schedule,
        present_count=present_count, late_count=late_count, absent_count=absent_count,
        attendance_percent=attendance_percent, face_profile=current_user.face_profile,
        finalized_total=finalized_total, now=datetime.now())


@student_bp.route("/courses", methods=["GET", "POST"])
@login_required
@role_required("student")
def courses():
    active = get_active_semester()
    selected_recommended_semester = request.args.get("recommended_semester", type=int)
    selected_type = request.args.get("course_type", "")
    selected_mode = request.args.get("delivery_mode", "")
    search = (request.args.get("q") or "").strip()
    enrollment_map = {}
    if active:
        rows = CourseEnrollment.query.join(ClassOffering).filter(
            CourseEnrollment.student_id == current_user.id,
            ClassOffering.semester_id == active.id,
        ).all()
        enrollment_map = {row.offering_id: row for row in rows}

    if request.method == "POST":
        offering = db.get_or_404(ClassOffering, request.form.get("offering_id", type=int))
        if not active or offering.semester_id != active.id:
            flash("Kelas bukan bagian dari semester aktif.", "danger")
        elif not active.enrollment_open:
            flash("Periode pengambilan mata kuliah sedang ditutup.", "warning")
        elif offering.status != "OPEN":
            flash("Penawaran kelas sudah ditutup.", "warning")
        elif offering.course.program_id != current_user.program_id:
            flash("Mata kuliah ini tidak tersedia untuk program studi Anda.", "danger")
        else:
            existing_same = CourseEnrollment.query.filter_by(
                student_id=current_user.id, offering_id=offering.id
            ).first()
            existing_course_request = CourseEnrollment.query.join(ClassOffering).filter(
                CourseEnrollment.student_id == current_user.id,
                CourseEnrollment.status.in_(("PENDING", "APPROVED")),
                ClassOffering.semester_id == active.id,
                ClassOffering.course_id == offering.course_id,
                ClassOffering.id != offering.id,
            ).first()
            if existing_same and existing_same.status != "REJECTED":
                flash("Anda sudah memiliki pengajuan atau enrolmen untuk kelas ini.", "info")
            elif existing_course_request:
                flash("Anda sudah mengajukan atau mengambil mata kuliah ini pada kelas lain di semester aktif.", "warning")
            else:
                active_requests = CourseEnrollment.query.filter(
                    CourseEnrollment.offering_id == offering.id,
                    CourseEnrollment.status.in_(("PENDING", "APPROVED")),
                ).count()
                if active_requests >= offering.capacity:
                    flash("Kelas sudah mencapai kapasitas maksimal termasuk pengajuan yang menunggu persetujuan.", "warning")
                else:
                    if existing_same and existing_same.status == "REJECTED":
                        enrollment = existing_same
                        enrollment.status = "PENDING"
                        enrollment.requested_at = datetime.now()
                        enrollment.reviewed_at = None
                        enrollment.reviewed_by_id = None
                        enrollment.review_note = None
                    else:
                        enrollment = CourseEnrollment(student_id=current_user.id, offering_id=offering.id, status="PENDING")
                        db.session.add(enrollment)
                    log_action("ENROLLMENT_REQUEST", f"Mahasiswa {current_user.npm} mengajukan {offering.course.code} {offering.class_code}")
                    db.session.commit()
                    flash("Pengajuan mata kuliah dikirim. Jadwal akan aktif setelah disetujui Admin.", "success")
                    return redirect(url_for("student.courses"))

    if active:
        from ..models import Course
        query = ClassOffering.query.join(Semester).join(Course).join(User, ClassOffering.lecturer_id == User.id).filter(
            ClassOffering.status == "OPEN",
            ClassOffering.semester_id == active.id,
            Course.is_active.is_(True),
            Course.program_id == current_user.program_id,
        )
        if selected_recommended_semester:
            query = query.filter(Course.recommended_semester == selected_recommended_semester)
        if selected_type in ("W", "P"):
            query = query.filter(Course.course_type == selected_type)
        if selected_mode in ("ONLINE", "OFFLINE"):
            query = query.filter(ClassOffering.schedules.any(ClassSchedule.delivery_mode == selected_mode))
        if search:
            term = f"%{search}%"
            query = query.filter(or_(
                Course.code.ilike(term), Course.name.ilike(term),
                ClassOffering.class_code.ilike(term), User.full_name.ilike(term)
            ))
        offerings = query.order_by(Course.recommended_semester, Course.code, ClassOffering.class_code).all()
    else:
        offerings = []

    enrollments = CourseEnrollment.query.join(ClassOffering).join(Semester).filter(
        CourseEnrollment.student_id == current_user.id
    ).order_by(CourseEnrollment.requested_at.desc()).all()
    programs = StudyProgram.query.filter_by(is_active=True).order_by(StudyProgram.name).all()
    return render_template(
        "student/courses.html", page_title="Mata Kuliah Saya", offerings=offerings,
        enrollment_map=enrollment_map, active_semester=active, enrollments=enrollments,
        selected_recommended_semester=selected_recommended_semester,
        selected_type=selected_type, selected_mode=selected_mode, search=search,
        semester_options=range(1, 9), programs=programs,
    )


@student_bp.route("/schedule")
@login_required
@role_required("student")
def schedule():
    enrollments = _active_enrollments()
    schedule_rows = []
    for enrollment in enrollments:
        for slot in enrollment.offering.schedules:
            schedule_rows.append((enrollment.offering, slot))
    schedule_rows.sort(key=lambda row: (row[1].weekday, row[1].start_time))
    active = get_active_semester()
    return render_template("student/schedule.html", page_title="Jadwal Kuliah", schedule_rows=schedule_rows, active_semester=active,
                           course_count=len(enrollments), total_credits=sum(e.offering.course.credits for e in enrollments))


@student_bp.route("/schedule/download.pdf")
@login_required
@role_required("student")
def download_schedule_pdf():
    enrollments = _active_enrollments()
    schedule_rows = [(enrollment.offering, slot) for enrollment in enrollments for slot in enrollment.offering.schedules]
    schedule_rows.sort(key=lambda row: (row[1].weekday, row[1].start_time))
    active = get_active_semester()
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
        from xml.sax.saxutils import escape
        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=26, leftMargin=26, topMargin=25, bottomMargin=25)
        styles = getSampleStyleSheet()
        brand_style = ParagraphStyle("ScheduleBrand", parent=styles["Title"], textColor=colors.HexColor("#173866"), alignment=TA_LEFT, fontSize=14, leading=17)
        subtitle_style = ParagraphStyle("ScheduleSub", parent=styles["Normal"], textColor=colors.HexColor("#63738b"), fontSize=8, leading=11)
        header_style = ParagraphStyle("ScheduleHeader", parent=styles["Normal"], textColor=colors.white, fontName="Helvetica-Bold", fontSize=7.5, leading=9, alignment=TA_CENTER)
        cell_style = ParagraphStyle("ScheduleCell", parent=styles["Normal"], textColor=colors.HexColor("#24364b"), fontSize=7.2, leading=9)
        title_cell = Paragraph("UNIVERSITAS TEKNOKRAT INDONESIA", brand_style)
        sub_cell = Paragraph("JADWAL MATA KULIAH MAHASISWA", subtitle_style)
        logo_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "images", "logo-teknokrat.png")
        try:
            logo = Image(logo_path, width=43, height=43) if os.path.exists(logo_path) else Paragraph("UTI", brand_style)
        except Exception:
            logo = Paragraph("UTI", brand_style)
        header = Table([[logo, [title_cell, sub_cell]]], colWidths=[55, 720])
        header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 2),
                                    ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story = [header, Spacer(1, 4)]
        semester_name = active.name if active else "Belum ada semester aktif"
        program_name = current_user.program.name if current_user.program else "-"
        meta = [["Nama", current_user.full_name, "NPM", current_user.npm or "-"],
                ["Program Studi", program_name, "Semester Akademik", semester_name]]
        meta_table = Table(meta, colWidths=[85, 245, 105, 340])
        meta_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf3fb")),
            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#edf3fb")),
            ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#173866")),
            ("TEXTCOLOR", (2, 0), (2, -1), colors.HexColor("#173866")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dbe3ee")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("PADDING", (0, 0), (-1, -1), 6), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story += [meta_table, Spacer(1, 12)]
        headers = ["Hari", "Jam", "Kode", "Mata Kuliah", "Kelas", "W/P", "SKS", "Dosen", "Moda", "Ruang / Link"]
        data = [[Paragraph(escape(item), header_style) for item in headers]]
        for offering, slot in schedule_rows:
            room_info = escape(slot.room or "-")
            if slot.delivery_mode == "ONLINE" and slot.online_url:
                room_info += "<br/><font color='#2456A6'>" + escape(slot.online_url) + "</font>"
            row = [slot.weekday_name, f"{slot.start_time:%H:%M}–{slot.end_time:%H:%M}", offering.course.code,
                   offering.course.name, offering.class_code, offering.course.course_type,
                   str(offering.course.credits), offering.lecturer.full_name,
                   slot.delivery_mode_label, room_info]
            data.append([Paragraph(escape(str(value)), cell_style) if index != 9 else Paragraph(str(value), cell_style)
                         for index, value in enumerate(row)])
        if len(data) == 1:
            data.append([Paragraph("Belum ada jadwal yang disetujui.", cell_style)] + [Paragraph("-", cell_style) for _ in headers[1:]])
        table = Table(data, repeatRows=1, colWidths=[43, 70, 52, 143, 48, 34, 32, 130, 58, 165], hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173866")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8e0eb")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fc")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        story += [table, Spacer(1, 12), Paragraph("Dokumen dibuat otomatis oleh Sistem Informasi Presensi Mahasiswa. Pastikan memeriksa portal untuk perubahan jadwal terbaru.", subtitle_style)]
        doc.build(story)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="jadwal-kuliah-semester.pdf", mimetype="application/pdf")
    except ImportError:
        flash("Library PDF belum terpasang. Jalankan: pip install -r requirements.txt", "danger")
        return redirect(url_for("student.schedule"))


@student_bp.route("/schedule/download.xlsx")
@login_required
@role_required("student")
def download_schedule_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage
    enrollments = _active_enrollments()
    schedule_rows = [(enrollment.offering, slot) for enrollment in enrollments for slot in enrollment.offering.schedules]
    schedule_rows.sort(key=lambda row: (row[1].weekday, row[1].start_time))
    active = get_active_semester()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jadwal Semester"
    sheet.merge_cells("B1:L1")
    sheet["B1"] = "UNIVERSITAS TEKNOKRAT INDONESIA"
    sheet["B1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    sheet["B1"].fill = PatternFill("solid", fgColor="173866")
    sheet["B1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 45
    logo_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "images", "logo-teknokrat.png")
    if os.path.exists(logo_path):
        try:
            logo = XLImage(logo_path)
            logo.height, logo.width = 38, 38
            sheet.add_image(logo, "A1")
        except Exception:
            pass
    sheet.merge_cells("A2:L2")
    sheet["A2"] = "JADWAL MATA KULIAH MAHASISWA"
    sheet["A2"].font = Font(size=12, bold=True, color="173866")
    sheet["A2"].alignment = Alignment(horizontal="center")
    sheet["A3"], sheet["B3"] = "Nama", current_user.full_name
    sheet["D3"], sheet["E3"] = "NPM", current_user.npm or "-"
    sheet["G3"], sheet["H3"] = "Program Studi", current_user.program.name if current_user.program else "-"
    sheet["J3"], sheet["K3"] = "Semester", active.name if active else "Belum ada semester aktif"
    sheet["A4"], sheet["B4"] = "Total Mata Kuliah", len(enrollments)
    sheet["D4"], sheet["E4"] = "Total SKS", sum(e.offering.course.credits for e in enrollments)
    headers = ["Hari", "Jam Mulai", "Jam Selesai", "Kode", "Mata Kuliah", "Kelas", "W/P", "SKS", "Dosen", "Moda", "Ruang", "Tautan Online"]
    header_row = 6
    for col, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2456A6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for offering, slot in schedule_rows:
        sheet.append([slot.weekday_name, slot.start_time.strftime("%H:%M"), slot.end_time.strftime("%H:%M"),
                      offering.course.code, offering.course.name, offering.class_code, offering.course.course_type,
                      offering.course.credits, offering.lecturer.full_name, slot.delivery_mode_label,
                      slot.room, slot.online_url or "-"])
    thin = Side(style="thin", color="DCE4EF")
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, min_col=1, max_col=12):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [13, 13, 13, 13, 36, 15, 8, 8, 27, 16, 16, 40]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A{header_row}:L{max(sheet.max_row, header_row)}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="jadwal-kuliah-semester.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@student_bp.route("/face-enrollment", methods=["GET"])
@login_required
@role_required("student")
def face_enrollment():
    """Render the student's face enrollment page.

    This endpoint is referenced by the student sidebar and dashboard. Keeping it
    separate from the API endpoint also means the camera UI can be loaded without
    submitting any biometric data.
    """
    return render_template(
        "student/face_enrollment.html",
        page_title="Pendaftaran Wajah",
        face_profile=current_user.face_profile,
    )


@student_bp.post("/api/face/enroll")
@login_required
@role_required("student")
def api_face_enroll():
    payload = request.get_json(silent=True) or {}
    frames = payload.get("frames") or {}
    image_data = frames.get("front") or payload.get("image")
    consent = payload.get("consent") is True
    if not consent:
        return jsonify(ok=False, message="Persetujuan pemrosesan data biometrik wajib diberikan sebelum pendaftaran."), 400
    if not image_data:
        return jsonify(ok=False, message="Sampel wajah depan tidak ditemukan."), 400
    try:
        # Server memeriksa sampel blink dan gerakan kepala sebelum menyimpan Face Profile.
        verify_enrollment_liveness(frames)
        embedding = extract_embedding(image_data)
        profile = current_user.face_profile
        now = datetime.now()
        if profile:
            profile.embedding_json = json.dumps(embedding)
            profile.model_name = "ArcFace"
            profile.consent_at = now
            profile.updated_at = now
        else:
            profile = FaceProfile(student_id=current_user.id, embedding_json=json.dumps(embedding),
                                  model_name="ArcFace", consent_at=now, created_at=now, updated_at=now)
            db.session.add(profile)
        log_action("FACE_ENROLLMENT", f"Mahasiswa {current_user.npm} melakukan pendaftaran/update Face Profile dengan liveness challenge")
        db.session.commit()
        return jsonify(ok=True, message="Face Profile berhasil diverifikasi dan disimpan. Challenge kedipan serta gerakan kepala lolos.")
    except FaceEngineError as exc:
        db.session.rollback()
        return jsonify(ok=False, message=str(exc)), 400
    except Exception:
        db.session.rollback()
        return jsonify(ok=False, message="Pendaftaran wajah gagal karena kesalahan sistem. Silakan coba kembali."), 500


@student_bp.route("/attendance")
@login_required
@role_required("student")
def attendance():
    enrollments = _active_enrollments()
    offering_ids = [e.offering_id for e in enrollments]
    sessions = AttendanceSession.query.filter(AttendanceSession.offering_id.in_(offering_ids)).order_by(
        AttendanceSession.opened_at.desc()).all() if offering_ids else []
    my_attendance = {a.session_id: a for a in Attendance.query.join(AttendanceSession).filter(
        Attendance.student_id == current_user.id, AttendanceSession.offering_id.in_(offering_ids)).all()} if offering_ids else {}
    token_value = (request.args.get("token") or "").strip()
    selected_session = None
    if token_value:
        from sqlalchemy import or_
        candidate = AttendanceSession.query.filter(or_(
            AttendanceSession.access_token == token_value,
            AttendanceSession.session_code == token_value.upper()
        )).first()
        now = datetime.now()
        if (candidate and candidate.offering_id in offering_ids and candidate.status == "ACTIVE"
                and candidate.opened_at <= now <= candidate.closes_at):
            selected_session = candidate
        else:
            flash("Token tidak valid, sesi sudah ditutup, atau sesi bukan untuk mata kuliah yang Anda ikuti.", "warning")
    now = datetime.now()
    return render_template("student/attendance.html", page_title="Presensi", sessions=sessions,
                           attendance_map=my_attendance, now=now, face_profile=current_user.face_profile,
                           selected_session_id=selected_session.id if selected_session else None,
                           token_input=token_value if selected_session else "")


@student_bp.post("/api/attendance/<int:session_id>/verify")
@login_required
@role_required("student")
def api_attendance_verify(session_id):
    payload = request.get_json(silent=True) or {}
    frames = payload.get("frames") or {}
    image_data = frames.get("open") or payload.get("image")
    token = (payload.get("token") or "").strip()
    if not token:
        return jsonify(ok=False, message="Token sesi wajib diisi. Scan QR dosen atau masukkan token sesi terlebih dahulu."), 400
    if not image_data:
        return jsonify(ok=False, message="Sampel wajah dengan mata terbuka tidak ditemukan."), 400
    if not current_user.active:
        return jsonify(ok=False, message="Akun tidak aktif."), 403
    profile = current_user.face_profile
    if not profile:
        return jsonify(ok=False, message="Face Profile belum terdaftar. Daftarkan wajah dari dashboard."), 400
    session = db.get_or_404(AttendanceSession, session_id)
    valid_token = (token == session.access_token) or (token.upper() == (session.session_code or "").upper())
    if not valid_token:
        return jsonify(ok=False, message="Token tidak cocok dengan sesi presensi yang dipilih."), 403
    enrollment = CourseEnrollment.query.filter_by(student_id=current_user.id, offering_id=session.offering_id, status="APPROVED").first()
    if not enrollment:
        return jsonify(ok=False, message="Anda tidak terdaftar sebagai peserta yang disetujui pada kelas ini."), 403
    now = datetime.now()
    if session.status != "ACTIVE" or now < session.opened_at or now > session.closes_at:
        return jsonify(ok=False, message="Sesi presensi belum dibuka atau sudah ditutup."), 400
    if Attendance.query.filter_by(session_id=session.id, student_id=current_user.id).first():
        return jsonify(ok=False, message="Anda sudah memiliki catatan presensi untuk sesi ini."), 409
    try:
        verify_blink_liveness(frames)
        match, score = compare_embedding(image_data, profile.embedding_json)
        if not match:
            log_action("FACE_VERIFICATION_FAILED", f"Mahasiswa {current_user.npm}; skor {score:.3f}")
            db.session.commit()
            return jsonify(ok=False, message="Wajah tidak cocok dengan Face Profile. Pastikan pencahayaan cukup dan hadapkan wajah ke kamera.", score=round(score, 3)), 401
        status = "PRESENT" if now <= session.late_after else "LATE"
        attendance_row = Attendance(session_id=session.id, student_id=current_user.id, status=status,
                                    checked_in_at=now, match_score=round(score, 5), verification_method="FACE_RECOGNITION_LIVENESS")
        db.session.add(attendance_row)
        log_action("ATTENDANCE_CHECK_IN", f"NPM {current_user.npm}, sesi #{session.session_number}, status {status}")
        db.session.commit()
        return jsonify(ok=True, message="Presensi berhasil dicatat setelah verifikasi kedipan dan wajah.", status=status,
                       status_label="Hadir" if status == "PRESENT" else "Terlambat",
                       time=now.strftime("%H:%M:%S"), score=round(score, 3))
    except FaceEngineError as exc:
        db.session.rollback()
        return jsonify(ok=False, message=str(exc)), 400
    except IntegrityError:
        db.session.rollback()
        return jsonify(ok=False, message="Presensi sudah tercatat pada sesi ini."), 409
    except Exception:
        db.session.rollback()
        return jsonify(ok=False, message="Verifikasi gagal karena kesalahan sistem. Silakan coba kembali."), 500


@student_bp.route("/history")
@login_required
@role_required("student")
def history():
    rows = Attendance.query.filter_by(student_id=current_user.id).join(AttendanceSession).order_by(AttendanceSession.opened_at.desc()).all()
    return render_template("student/history.html", page_title="Riwayat Presensi", attendances=rows)


@student_bp.route("/summary")
@login_required
@role_required("student")
def summary():
    enrollments = _active_enrollments()
    data = []
    for enrollment in enrollments:
        sessions = AttendanceSession.query.filter_by(offering_id=enrollment.offering_id, status="CLOSED").all()
        session_ids = [s.id for s in sessions]
        records = Attendance.query.filter(Attendance.student_id == current_user.id, Attendance.session_id.in_(session_ids)).all() if session_ids else []
        present = sum(1 for a in records if a.status == "PRESENT")
        late = sum(1 for a in records if a.status == "LATE")
        absent = sum(1 for a in records if a.status == "ABSENT")
        total = len(sessions)
        percent = round((present + late) * 100 / total) if total else 0
        data.append({"offering": enrollment.offering, "total": total, "present": present, "late": late, "absent": absent, "percent": percent})
    return render_template("student/summary.html", page_title="Rekap Kehadiran", summary=data)


@student_bp.route("/notifications")
@login_required
@role_required("student")
def notifications():
    items = Notification.query.filter_by(user_id=current_user.id).order_by(Notification.created_at.desc()).limit(100).all()
    for item in items:
        item.is_read = True
    db.session.commit()
    return render_template("student/notifications.html", page_title="Notifikasi", notifications=items)


@student_bp.route("/profile", methods=["GET", "POST"])
@login_required
@role_required("student")
def profile():
    from sqlalchemy import func
    if request.method == "POST":
        full_name = (request.form.get("full_name") or "").strip()
        email = (request.form.get("email") or "").strip().lower()
        current_password = request.form.get("current_password") or ""
        new_password = request.form.get("new_password") or ""
        if not full_name or not email:
            flash("Nama dan email wajib diisi.", "danger")
        elif User.query.filter(func.lower(User.email) == email, User.id != current_user.id).first():
            flash("Email sudah digunakan oleh akun lain.", "danger")
        elif new_password and (not current_user.check_password(current_password) or len(new_password) < 8):
            flash("Untuk mengganti password, masukkan password sekarang dengan benar dan password baru minimal 8 karakter.", "danger")
        else:
            current_user.full_name = full_name
            current_user.email = email
            if new_password:
                current_user.set_password(new_password)
            log_action("STUDENT_PROFILE_UPDATE", "Mahasiswa memperbarui profil akun")
            db.session.commit()
            flash("Profil berhasil diperbarui.", "success")
            return redirect(url_for("student.profile"))
    return render_template("student/profile.html", page_title="Profil Saya")
