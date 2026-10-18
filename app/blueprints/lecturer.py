from datetime import date, datetime, time, timedelta
from io import BytesIO
import base64
import json
import os
import secrets
import string
from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import login_required, current_user
from sqlalchemy.exc import IntegrityError
from ..extensions import db
from ..models import (User, StudyProgram, Semester, Course, ClassOffering, ClassSchedule,
                      CourseEnrollment, AttendanceSession, Attendance, Notification)
from ..utils import role_required, log_action, get_active_semester

lecturer_bp = Blueprint("lecturer", __name__, url_prefix="/lecturer")
WEEKDAYS = [(0, "Senin"), (1, "Selasa"), (2, "Rabu"), (3, "Kamis"), (4, "Jumat"), (5, "Sabtu"), (6, "Minggu")]
WEEKDAY_NAMES = [label for _, label in WEEKDAYS]


def _owned_offering(offering_id):
    return ClassOffering.query.filter_by(id=offering_id, lecturer_id=current_user.id).first_or_404()


def _owned_session(session_id):
    session = db.get_or_404(AttendanceSession, session_id)
    if session.offering.lecturer_id != current_user.id:
        abort(403)
    return session


def _parse_time(value):
    return datetime.strptime(value, "%H:%M").time()


def _slots_from_form(delivery_mode="OFFLINE"):
    weekdays = request.form.getlist("weekday[]")
    starts = request.form.getlist("start_time[]")
    ends = request.form.getlist("end_time[]")
    rooms = request.form.getlist("room[]")
    modes = request.form.getlist("delivery_mode[]")
    urls = request.form.getlist("online_url[]")
    legacy_mode = (request.form.get("delivery_mode") or delivery_mode or "OFFLINE").upper()
    legacy_url = (request.form.get("online_url") or "").strip() or None
    slots = []
    count = max(len(weekdays), len(starts), len(ends), len(rooms), len(modes), len(urls))
    for i in range(count):
        wd = weekdays[i] if i < len(weekdays) else ""
        start = starts[i] if i < len(starts) else ""
        end = ends[i] if i < len(ends) else ""
        room = rooms[i] if i < len(rooms) else ""
        if not str(wd).strip() and not str(start).strip() and not str(end).strip():
            continue
        mode = (modes[i] if i < len(modes) else legacy_mode).upper()
        online_url = (urls[i] if i < len(urls) else legacy_url) or ""
        online_url = online_url.strip() or None
        if mode not in ("ONLINE", "OFFLINE"):
            raise ValueError("Pilih moda online atau tatap muka pada setiap slot.")
        if mode == "ONLINE" and online_url and not online_url.lower().startswith(("https://", "http://")):
            raise ValueError("Tautan meeting online harus dimulai dengan http:// atau https://.")
        try:
            weekday = int(wd)
            start_time, end_time = _parse_time(start), _parse_time(end)
        except (ValueError, TypeError):
            raise ValueError("Setiap slot jadwal wajib memiliki hari, jam mulai, dan jam selesai yang valid.")
        room = (room or "").strip()
        if mode == "ONLINE":
            room = "Online"
        if weekday not in range(7) or start_time >= end_time:
            raise ValueError("Hari atau rentang waktu jadwal tidak valid.")
        if mode == "OFFLINE" and not room:
            raise ValueError("Ruang wajib diisi untuk kelas tatap muka.")
        slots.append({"weekday": weekday, "start_time": start_time, "end_time": end_time,
                      "room": room or "Online", "delivery_mode": mode,
                      "online_url": online_url if mode == "ONLINE" else None})
    if not slots:
        raise ValueError("Tambahkan minimal satu jadwal mingguan.")
    for i, one in enumerate(slots):
        for other in slots[i + 1:]:
            if one["weekday"] == other["weekday"] and one["start_time"] < other["end_time"] and other["start_time"] < one["end_time"]:
                raise ValueError("Dua slot jadwal yang dibuat saling bertabrakan.")
    return slots


def _aggregate_delivery_mode(slots):
    modes = {slot["delivery_mode"] for slot in slots}
    return next(iter(modes)) if len(modes) == 1 else "MIXED"


def _legacy_online_url(slots):
    return next((slot["online_url"] for slot in slots if slot["delivery_mode"] == "ONLINE" and slot.get("online_url")), None)


def _validate_conflicts(semester_id, lecturer_id, slots, delivery_mode=None, exclude_offering_id=None):
    query = ClassSchedule.query.join(ClassOffering).filter(ClassOffering.semester_id == semester_id)
    if exclude_offering_id:
        query = query.filter(ClassOffering.id != exclude_offering_id)
    existing_slots = query.all()
    for slot in slots:
        for existing in existing_slots:
            if existing.weekday != slot["weekday"]:
                continue
            overlaps = slot["start_time"] < existing.end_time and existing.start_time < slot["end_time"]
            if not overlaps:
                continue
            if existing.offering.lecturer_id == lecturer_id:
                raise ValueError(f"Jadwal bentrok: Anda sudah mengajar {existing.offering.course.code} {existing.offering.class_code} pada waktu tersebut.")
            existing_mode = getattr(existing, "delivery_mode", None) or existing.offering.delivery_mode or "OFFLINE"
            if slot["delivery_mode"] == "OFFLINE" and existing_mode == "OFFLINE" and existing.room.casefold() == slot["room"].casefold():
                raise ValueError(f"Ruang {slot['room']} sudah digunakan oleh kelas {existing.offering.course.code} {existing.offering.class_code} pada waktu tersebut.")


def _make_session_qr(session):
    if not session.access_token:
        return None, None
    root = (os.getenv("APP_BASE_URL") or request.host_url).rstrip("/")
    qr_url = root + url_for("student.attendance", token=session.access_token)
    try:
        import qrcode
        qr = qrcode.QRCode(version=2, box_size=7, border=3)
        qr.add_data(qr_url)
        qr.make(fit=True)
        image = qr.make_image(fill_color="#142f56", back_color="white")
        output = BytesIO()
        image.save(output, format="PNG")
        data_uri = "data:image/png;base64," + base64.b64encode(output.getvalue()).decode("ascii")
        return data_uri, qr_url
    except ImportError:
        return None, qr_url


def _attendance_rows(sessions):
    session_ids = [session.id for session in sessions]
    records = Attendance.query.filter(Attendance.session_id.in_(session_ids)).all() if session_ids else []
    return records


def _period_bounds(period):
    now = datetime.now()
    if period == "day":
        start = datetime.combine(now.date(), time.min)
        return start, start + timedelta(days=1), "Harian"
    if period == "month":
        start = datetime(now.year, now.month, 1)
        end = datetime(now.year + (now.month == 12), 1 if now.month == 12 else now.month + 1, 1)
        return start, end, "Bulanan"
    active = get_active_semester()
    if active and active.start_date:
        start = datetime.combine(active.start_date, time.min)
        end = datetime.combine(active.end_date + timedelta(days=1), time.min) if active.end_date else datetime.max
        return start, end, f"Semester · {active.name}"
    return None, None, "Semua Periode"


def _build_xlsx(title, metadata, headers, rows, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Laporan"
    column_count = max(len(headers), 2)
    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=column_count)
    title_cell = sheet.cell(1, 2, "UNIVERSITAS TEKNOKRAT INDONESIA")
    title_cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="173866")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(2, column_count + 1):
        sheet.cell(1, col).fill = PatternFill("solid", fgColor="173866")
    sheet.row_dimensions[1].height = 42
    logo_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "images", "logo-teknokrat.png")
    if os.path.exists(logo_path):
        try:
            logo = XLImage(logo_path)
            logo.height, logo.width = 35, 35
            sheet.add_image(logo, "A1")
        except Exception:
            pass
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=column_count)
    subtitle = sheet.cell(2, 1, title.upper())
    subtitle.font = Font(size=12, bold=True, color="173866")
    subtitle.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 24
    for item in metadata:
        row_num = sheet.max_row + 1
        key = sheet.cell(row_num, 1, str(item[0]))
        key.font = Font(bold=True, color="173866")
        key.fill = PatternFill("solid", fgColor="EDF3FB")
        value = sheet.cell(row_num, 2, str(item[1]))
        value.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.append([])
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2456A6")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        sheet.append([value if value is not None else "-" for value in row])
    thin = Side(style="thin", color="DCE4EF")
    for row in sheet.iter_rows(min_row=header_row, max_row=sheet.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx in range(1, len(headers) + 1):
        values = [len(str(sheet.cell(r, idx).value or "")) for r in range(header_row, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(max(values or [12]) + 2, 12), 38)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _build_pdf(title, metadata, headers, rows, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    from xml.sax.saxutils import escape

    output = BytesIO()
    page_width = landscape(A4)[0] - 52
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=26, leftMargin=26, topMargin=25, bottomMargin=25)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("UTITitle", parent=styles["Title"], textColor=colors.HexColor("#173866"), alignment=TA_LEFT, fontSize=15, leading=18)
    subtitle_style = ParagraphStyle("UTISubtitle", parent=styles["Normal"], textColor=colors.HexColor("#63738b"), fontSize=8, leading=11)
    header_style = ParagraphStyle("UTITableHeader", parent=styles["Normal"], textColor=colors.white, fontName="Helvetica-Bold", fontSize=7, leading=8, alignment=TA_CENTER)
    cell_style = ParagraphStyle("UTITableCell", parent=styles["Normal"], textColor=colors.HexColor("#25364d"), fontSize=7, leading=8.5, wordWrap="CJK")
    logo_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "images", "logo-teknokrat.png")
    try:
        logo = Image(logo_path, width=42, height=42) if os.path.exists(logo_path) else Paragraph("UTI", title_style)
    except Exception:
        logo = Paragraph("UTI", title_style)
    brand = [Paragraph("UNIVERSITAS TEKNOKRAT INDONESIA", title_style),
             Paragraph(escape(title.upper()), subtitle_style)]
    header = Table([[logo, brand]], colWidths=[52, page_width - 52])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story = [header, Spacer(1, 5)]
    meta_rows = [[Paragraph(f"<b>{escape(str(k))}</b>", cell_style), Paragraph(escape(str(v)), cell_style)] for k, v in metadata]
    if meta_rows:
        meta_table = Table(meta_rows, colWidths=[110, page_width - 110])
        meta_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf3fb")),
                                        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#173866")),
                                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dbe3ee")),
                                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                        ("PADDING", (0, 0), (-1, -1), 6)]))
        story += [meta_table, Spacer(1, 10)]
    data = [[Paragraph(escape(str(value)), header_style) for value in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(value if value is not None else "-")), cell_style) for value in row])
    if not rows:
        data.append([Paragraph("Tidak ada data untuk filter yang dipilih.", cell_style)] + [Paragraph("-", cell_style) for _ in headers[1:]])
    col_width = page_width / max(len(headers), 1)
    table = Table(data, repeatRows=1, colWidths=[col_width] * len(headers))
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173866")),
                               ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d8e0eb")),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fc")]),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                               ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.append(table)
    story.append(Spacer(1, 10))
    story.append(Paragraph("Dokumen dibuat otomatis oleh Sistem Informasi Presensi Mahasiswa Universitas Teknokrat Indonesia.", subtitle_style))
    doc.build(story)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/pdf")


@lecturer_bp.route("/")
@login_required
@role_required("lecturer")
def dashboard():
    offerings = ClassOffering.query.filter_by(lecturer_id=current_user.id).order_by(ClassOffering.id.desc()).all()
    offering_ids = [o.id for o in offerings]
    sessions = AttendanceSession.query.filter(AttendanceSession.offering_id.in_(offering_ids)).order_by(AttendanceSession.opened_at.desc()).limit(8).all() if offering_ids else []
    active_count = sum(1 for s in sessions if s.status == "ACTIVE" and s.closes_at > datetime.now())
    todays = []
    weekday = datetime.now().weekday()
    for offering in offerings:
        for schedule in offering.schedules:
            if schedule.weekday == weekday:
                todays.append((offering, schedule))
    schedule_rows = sorted(todays, key=lambda row: row[1].start_time)
    records = _attendance_rows(sessions)
    stat_counts = {"PRESENT": sum(1 for a in records if a.status == "PRESENT"),
                   "LATE": sum(1 for a in records if a.status == "LATE"),
                   "ABSENT": sum(1 for a in records if a.status == "ABSENT")}
    return render_template("lecturer/dashboard.html", page_title="Dashboard Dosen", offerings=offerings,
                           sessions=sessions, active_count=active_count, todays=schedule_rows,
                           stat_counts=stat_counts, now=datetime.now())


@lecturer_bp.route("/classes")
@login_required
@role_required("lecturer")
def classes():
    offerings = ClassOffering.query.filter_by(lecturer_id=current_user.id).order_by(ClassOffering.id.desc()).all()
    return render_template("lecturer/classes.html", page_title="Kelas yang Diampu", offerings=offerings)


@lecturer_bp.route("/teaching/add", methods=["GET", "POST"])
@login_required
@role_required("lecturer")
def add_teaching():
    programs = StudyProgram.query.filter_by(is_active=True).order_by(StudyProgram.name).all()
    courses = Course.query.filter_by(is_active=True).order_by(Course.program_id, Course.recommended_semester, Course.code).all()
    semesters = Semester.query.order_by(Semester.id.desc()).all()
    if request.method == "POST":
        mode = (request.form.get("course_entry_mode") or "catalog").lower()
        delivery_mode = (request.form.get("delivery_mode") or "OFFLINE").upper()
        semester = db.session.get(Semester, request.form.get("semester_id", type=int))
        class_code = (request.form.get("class_code") or "").strip().upper()
        capacity = request.form.get("capacity", type=int) or 40
        online_url = (request.form.get("online_url") or "").strip() or None
        course = None
        try:
            if mode == "new":
                code = (request.form.get("new_course_code") or "").strip().upper()
                name = (request.form.get("new_course_name") or "").strip()
                program = db.session.get(StudyProgram, request.form.get("new_program_id", type=int))
                credits = request.form.get("new_credits", type=int) or 2
                rec_sem = request.form.get("new_recommended_semester", type=int)
                course_type = (request.form.get("new_course_type") or "W").upper()
                existing_course = Course.query.filter_by(code=code).first() if code else None
                if existing_course:
                    raise ValueError("Kode mata kuliah sudah ada di katalog. Pilih mata kuliah tersebut dari daftar atau gunakan kode baru.")
                if (not code or not name or not program or credits not in range(1, 7)
                        or rec_sem not in range(1, 9) or course_type not in ("W", "P")):
                    raise ValueError("Lengkapi kode, nama, program studi, semester kurikulum 1–8, SKS, dan jenis mata kuliah baru.")
                course = Course(code=code, name=name, program_id=program.id, credits=credits,
                                recommended_semester=rec_sem, course_type=course_type,
                                description=(request.form.get("new_description") or "").strip() or None,
                                is_active=True)
                db.session.add(course)
                db.session.flush()
            else:
                course = db.session.get(Course, request.form.get("course_id", type=int))
                if not course or not course.is_active:
                    raise ValueError("Pilih mata kuliah aktif dari katalog.")
            if not semester or not semester.is_active:
                raise ValueError("Pilih semester akademik yang aktif.")
            if not class_code or capacity < 1:
                raise ValueError("Kode kelas dan kapasitas harus valid.")
            slots = _slots_from_form(delivery_mode)
            _validate_conflicts(semester.id, current_user.id, slots)
            delivery_mode = _aggregate_delivery_mode(slots)
            online_url = _legacy_online_url(slots)
            offering = ClassOffering(course=course, semester=semester, lecturer=current_user,
                                     class_code=class_code, capacity=capacity, status="OPEN",
                                     delivery_mode=delivery_mode, online_url=online_url,
                                     created_by_id=current_user.id)
            offering.schedules.extend(ClassSchedule(**slot) for slot in slots)
            db.session.add(offering)
            log_action("LECTURER_OFFERING_CREATE", f"Dosen membuat {course.code} kelas {class_code}; {semester.name}")
            db.session.commit()
            flash("Mata kuliah dan jadwal mengajar berhasil ditambahkan. Kelas tersedia untuk pengambilan mahasiswa.", "success")
            return redirect(url_for("lecturer.classes"))
        except (ValueError, IntegrityError) as exc:
            db.session.rollback()
            flash(str(exc) if isinstance(exc, ValueError) else "Kode kelas sudah digunakan untuk mata kuliah dan semester ini.", "danger")
    return render_template("lecturer/teaching_add.html", page_title="Input Mata Kuliah", programs=programs,
                           courses=courses, semesters=semesters, weekdays=WEEKDAYS)


@lecturer_bp.route("/classes/<int:offering_id>", methods=["GET", "POST"])
@login_required
@role_required("lecturer")
def class_detail(offering_id):
    offering = _owned_offering(offering_id)
    if request.method == "POST":
        session_number = request.form.get("session_number", type=int)
        title = (request.form.get("title") or "").strip() or f"Pertemuan {session_number}"
        open_minutes = request.form.get("open_minutes", type=int)
        grace_minutes = request.form.get("grace_minutes", type=int)
        if open_minutes is None:
            open_minutes = 10
        if grace_minutes is None:
            grace_minutes = 5
        if offering.status != "OPEN":
            flash("Kelas sudah ditutup. Sesi baru tidak dapat dibuka.", "danger")
        elif not session_number or session_number < 1 or open_minutes < 1 or open_minutes > 180 or grace_minutes < 0 or grace_minutes > 120:
            flash("Nomor pertemuan atau durasi presensi tidak valid.", "danger")
        elif AttendanceSession.query.filter_by(offering_id=offering.id, session_number=session_number).first():
            flash("Nomor pertemuan tersebut sudah digunakan.", "danger")
        else:
            now = datetime.now()
            token = secrets.token_urlsafe(24)
            alphabet = string.ascii_uppercase + string.digits
            code = "".join(secrets.choice(alphabet) for _ in range(6))
            while AttendanceSession.query.filter_by(session_code=code).first():
                code = "".join(secrets.choice(alphabet) for _ in range(6))
            late_after = now + timedelta(minutes=open_minutes)
            closes_at = late_after + timedelta(minutes=grace_minutes)
            session = AttendanceSession(offering_id=offering.id, session_number=session_number,
                title=title, status="ACTIVE", opened_at=now, late_after=late_after,
                closes_at=closes_at, opened_by_id=current_user.id,
                access_token=token, session_code=code)
            db.session.add(session)
            log_action("ATTENDANCE_SESSION_OPEN", f"Membuka sesi {session_number} {offering.course.code} {offering.class_code}; kode {code}")
            db.session.commit()
            flash("Sesi presensi dibuka. Tampilkan QR untuk smartphone atau kode token untuk laptop/komputer.", "success")
            return redirect(url_for("lecturer.session_detail", session_id=session.id))
    enrollments = CourseEnrollment.query.filter_by(offering_id=offering.id, status="APPROVED").all()
    enrollments.sort(key=lambda item: item.student.full_name.lower())
    sessions = AttendanceSession.query.filter_by(offering_id=offering.id).order_by(AttendanceSession.session_number.desc()).all()
    return render_template("lecturer/class_detail.html", page_title=offering.course.name, offering=offering,
                           enrollments=enrollments, sessions=sessions, now=datetime.now())


@lecturer_bp.route("/sessions/<int:session_id>")
@login_required
@role_required("lecturer")
def session_detail(session_id):
    session = _owned_session(session_id)
    enrolled_ids = [e.student_id for e in CourseEnrollment.query.filter_by(offering_id=session.offering_id, status="APPROVED").all()]
    existing = {a.student_id: a for a in Attendance.query.filter_by(session_id=session.id).all()}
    rows = []
    for student_id in enrolled_ids:
        student = db.session.get(User, student_id)
        rows.append((student, existing.get(student_id)))
    rows.sort(key=lambda item: item[0].full_name if item[0] else "")
    qr_data_uri, qr_url = _make_session_qr(session)
    return render_template("lecturer/session_detail.html", page_title="Monitoring Presensi", session=session,
                           rows=rows, now=datetime.now(), qr_data_uri=qr_data_uri, qr_url=qr_url)


@lecturer_bp.post("/sessions/<int:session_id>/close")
@login_required
@role_required("lecturer")
def close_session(session_id):
    session = _owned_session(session_id)
    if session.status == "CLOSED":
        flash("Sesi sudah ditutup sebelumnya.", "info")
        return redirect(url_for("lecturer.session_detail", session_id=session.id))
    now = datetime.now()
    approved = CourseEnrollment.query.filter_by(offering_id=session.offering_id, status="APPROVED").all()
    existing = {a.student_id for a in Attendance.query.filter_by(session_id=session.id).all()}
    try:
        for enrollment in approved:
            if enrollment.student_id not in existing:
                db.session.add(Attendance(session_id=session.id, student_id=enrollment.student_id,
                    status="ABSENT", checked_in_at=None, verification_method="SESSION_FINALIZATION"))
        session.status = "CLOSED"
        session.closed_at = now
        log_action("ATTENDANCE_SESSION_CLOSE", f"Menutup sesi #{session.session_number} kelas {session.offering.class_code}; finalisasi {len(approved) - len(existing)} data")
        db.session.commit()
        flash("Sesi ditutup dan rekap ketidakhadiran telah difinalisasi.", "success")
    except Exception:
        db.session.rollback()
        flash("Sesi belum berhasil ditutup. Periksa log aplikasi dan coba lagi.", "danger")
    return redirect(url_for("lecturer.session_detail", session_id=session.id))


@lecturer_bp.post("/attendance/<int:attendance_id>/correct")
@login_required
@role_required("lecturer")
def correct_attendance(attendance_id):
    attendance = db.get_or_404(Attendance, attendance_id)
    if attendance.session.offering.lecturer_id != current_user.id:
        abort(403)
    new_status = (request.form.get("status") or "").upper()
    reason = (request.form.get("reason") or "").strip()
    if new_status not in ("PRESENT", "LATE", "ABSENT") or not reason:
        flash("Status dan alasan koreksi wajib diisi.", "danger")
    elif attendance.session.status != "CLOSED":
        flash("Koreksi manual hanya diperbolehkan setelah sesi ditutup.", "warning")
    else:
        attendance.status = new_status
        attendance.correction_reason = reason[:255]
        attendance.corrected_by_id = current_user.id
        if new_status in ("PRESENT", "LATE") and attendance.checked_in_at is None:
            attendance.checked_in_at = datetime.now()
        log_action("ATTENDANCE_MANUAL_CORRECTION", f"Koreksi presensi #{attendance.id} menjadi {new_status}; alasan: {reason}")
        db.session.commit()
        flash("Koreksi presensi berhasil dicatat di log aktivitas.", "success")
    return redirect(url_for("lecturer.session_detail", session_id=attendance.session_id))


@lecturer_bp.route("/schedule")
@login_required
@role_required("lecturer")
def schedule():
    semesters = Semester.query.order_by(Semester.id.desc()).all()
    selected_id = request.args.get("semester_id", type=int)
    selected = db.session.get(Semester, selected_id) if selected_id else get_active_semester()
    query = ClassOffering.query.filter_by(lecturer_id=current_user.id)
    if selected:
        query = query.filter_by(semester_id=selected.id)
    offerings = query.join(ClassOffering.course).order_by(Course.code, ClassOffering.class_code).all()
    rows = [(offering, slot) for offering in offerings for slot in offering.schedules]
    rows.sort(key=lambda row: (row[1].weekday, row[1].start_time))
    return render_template("lecturer/schedule.html", page_title="Jadwal Mengajar", offerings=offerings,
                           schedule_rows=rows, semesters=semesters, selected_semester=selected)


@lecturer_bp.route("/schedule/download.pdf")
@login_required
@role_required("lecturer")
def download_schedule_pdf():
    selected_id = request.args.get("semester_id", type=int)
    selected = db.session.get(Semester, selected_id) if selected_id else get_active_semester()
    query = ClassOffering.query.filter_by(lecturer_id=current_user.id)
    if selected:
        query = query.filter_by(semester_id=selected.id)
    offerings = query.join(ClassOffering.course).order_by(Course.code).all()
    rows = [(o, s) for o in offerings for s in o.schedules]
    rows.sort(key=lambda row: (row[1].weekday, row[1].start_time))
    data = [[slot.weekday_name, f"{slot.start_time:%H:%M}–{slot.end_time:%H:%M}", offering.course.code,
             offering.course.name, offering.class_code, f"{offering.course.credits} SKS",
             slot.delivery_mode_label, offering.lecturer.full_name,
             (slot.room + (" · " + slot.online_url if slot.delivery_mode == "ONLINE" and slot.online_url else ""))]
            for offering, slot in rows]
    return _build_pdf("Jadwal Mengajar Dosen", [("Dosen", current_user.full_name), ("Semester", selected.name if selected else "Semua Semester")],
                      ["Hari", "Jam", "Kode", "Mata Kuliah", "Kelas", "SKS", "Moda", "Dosen", "Ruang/Online"], data,
                      "jadwal-mengajar.pdf")


@lecturer_bp.route("/schedule/download.xlsx")
@login_required
@role_required("lecturer")
def download_schedule_xlsx():
    selected_id = request.args.get("semester_id", type=int)
    selected = db.session.get(Semester, selected_id) if selected_id else get_active_semester()
    query = ClassOffering.query.filter_by(lecturer_id=current_user.id)
    if selected:
        query = query.filter_by(semester_id=selected.id)
    offerings = query.join(ClassOffering.course).order_by(Course.code).all()
    rows = []
    for offering in offerings:
        for slot in offering.schedules:
            rows.append([slot.weekday_name, slot.start_time.strftime("%H:%M"), slot.end_time.strftime("%H:%M"),
                offering.course.code, offering.course.name, offering.class_code, offering.course.credits,
                offering.course.course_type, slot.delivery_mode_label, offering.lecturer.full_name,
                slot.room, slot.online_url or "-"])
    rows.sort(key=lambda row: (WEEKDAY_NAMES.index(row[0]), row[1]))
    return _build_xlsx("Jadwal Mengajar Dosen", [("Dosen", current_user.full_name), ("Semester", selected.name if selected else "Semua Semester")],
                       ["Hari", "Mulai", "Selesai", "Kode", "Mata Kuliah", "Kelas", "SKS", "W/P", "Moda", "Dosen", "Ruang", "Tautan Online"],
                       rows, "jadwal-mengajar.xlsx")


@lecturer_bp.route("/statistics")
@login_required
@role_required("lecturer")
def statistics():
    period = request.args.get("period", "semester")
    if period not in ("day", "month", "semester"):
        period = "semester"
    start, end, period_label = _period_bounds(period)
    offering_id = request.args.get("offering_id", type=int)
    offering_query = ClassOffering.query.filter_by(lecturer_id=current_user.id)
    if period == "semester":
        active = get_active_semester()
        if active:
            offering_query = offering_query.filter_by(semester_id=active.id)
    if offering_id:
        offering_query = offering_query.filter_by(id=offering_id)
    offerings = offering_query.join(ClassOffering.course).order_by(Course.code).all()
    offering_ids = [o.id for o in offerings]
    session_query = AttendanceSession.query.filter(AttendanceSession.offering_id.in_(offering_ids)) if offering_ids else AttendanceSession.query.filter(AttendanceSession.id == -1)
    if start:
        session_query = session_query.filter(AttendanceSession.opened_at >= start)
    if end:
        session_query = session_query.filter(AttendanceSession.opened_at < end)
    sessions = session_query.order_by(AttendanceSession.opened_at.desc()).all()
    records = _attendance_rows(sessions)
    counts = {"PRESENT": sum(1 for row in records if row.status == "PRESENT"),
              "LATE": sum(1 for row in records if row.status == "LATE"),
              "ABSENT": sum(1 for row in records if row.status == "ABSENT")}
    by_offering = []
    sessions_by_offering = {offering_id: [] for offering_id in offering_ids}
    for session in sessions:
        sessions_by_offering.setdefault(session.offering_id, []).append(session)
    for offering in offerings:
        relevant_sessions = sessions_by_offering.get(offering.id, [])
        relevant_ids = {s.id for s in relevant_sessions}
        relevant = [a for a in records if a.session_id in relevant_ids]
        by_offering.append({"offering": offering, "sessions": len(relevant_sessions),
            "present": sum(1 for a in relevant if a.status == "PRESENT"),
            "late": sum(1 for a in relevant if a.status == "LATE"),
            "absent": sum(1 for a in relevant if a.status == "ABSENT")})
    chart_payload = json.dumps({"labels": [row["offering"].course.code + " " + row["offering"].class_code for row in by_offering],
        "present": [row["present"] for row in by_offering], "late": [row["late"] for row in by_offering],
        "absent": [row["absent"] for row in by_offering]})
    return render_template("lecturer/statistics.html", page_title="Statistik Kehadiran", offerings=offerings,
        by_offering=by_offering, counts=counts, sessions=sessions, period=period, period_label=period_label,
        selected_offering_id=offering_id, chart_payload=chart_payload)


@lecturer_bp.route("/statistics/export.xlsx")
@login_required
@role_required("lecturer")
def export_statistics_xlsx():
    period = request.args.get("period", "semester")
    if period not in ("day", "month", "semester"):
        period = "semester"
    start, end, label = _period_bounds(period)
    offering_id = request.args.get("offering_id", type=int)
    query = ClassOffering.query.filter_by(lecturer_id=current_user.id)
    if period == "semester" and get_active_semester():
        query = query.filter_by(semester_id=get_active_semester().id)
    if offering_id:
        query = query.filter_by(id=offering_id)
    offerings = query.all()
    ids = [o.id for o in offerings]
    sq = AttendanceSession.query.filter(AttendanceSession.offering_id.in_(ids)) if ids else AttendanceSession.query.filter(AttendanceSession.id == -1)
    if start:
        sq = sq.filter(AttendanceSession.opened_at >= start)
    if end:
        sq = sq.filter(AttendanceSession.opened_at < end)
    sessions = sq.order_by(AttendanceSession.opened_at.desc()).all()
    session_map = {s.id: s for s in sessions}
    rows = []
    for attendance in _attendance_rows(sessions):
        session = session_map.get(attendance.session_id)
        rows.append([session.offering.course.code, session.offering.course.name, session.offering.class_code,
            session.session_number, session.opened_at.strftime("%d/%m/%Y %H:%M"), attendance.student.npm,
            attendance.student.full_name, {"PRESENT":"Hadir", "LATE":"Terlambat", "ABSENT":"Tidak Hadir"}.get(attendance.status, attendance.status),
            attendance.checked_in_at.strftime("%H:%M:%S") if attendance.checked_in_at else "-"])
    rows.sort(key=lambda row: (row[0], row[4], row[5]))
    return _build_xlsx("Statistik dan Rekap Presensi Dosen", [("Dosen", current_user.full_name), ("Periode", label)],
        ["Kode", "Mata Kuliah", "Kelas", "Pertemuan", "Tanggal Sesi", "NPM", "Nama Mahasiswa", "Status", "Waktu Hadir"], rows,
        "statistik-kehadiran.xlsx")


@lecturer_bp.route("/statistics/export.pdf")
@login_required
@role_required("lecturer")
def export_statistics_pdf():
    period = request.args.get("period", "semester")
    if period not in ("day", "month", "semester"):
        period = "semester"
    start, end, label = _period_bounds(period)
    offering_id = request.args.get("offering_id", type=int)
    query = ClassOffering.query.filter_by(lecturer_id=current_user.id)
    if period == "semester" and get_active_semester():
        query = query.filter_by(semester_id=get_active_semester().id)
    if offering_id:
        query = query.filter_by(id=offering_id)
    offerings = query.all()
    ids = [o.id for o in offerings]
    sq = AttendanceSession.query.filter(AttendanceSession.offering_id.in_(ids)) if ids else AttendanceSession.query.filter(AttendanceSession.id == -1)
    if start:
        sq = sq.filter(AttendanceSession.opened_at >= start)
    if end:
        sq = sq.filter(AttendanceSession.opened_at < end)
    sessions = sq.order_by(AttendanceSession.opened_at.desc()).all()
    session_map = {s.id: s for s in sessions}
    rows = []
    for attendance in _attendance_rows(sessions):
        session = session_map.get(attendance.session_id)
        rows.append([session.offering.course.code, session.offering.course.name, session.offering.class_code,
            str(session.session_number), session.opened_at.strftime("%d/%m/%Y"), attendance.student.npm,
            attendance.student.full_name, {"PRESENT":"Hadir", "LATE":"Terlambat", "ABSENT":"Tidak Hadir"}.get(attendance.status, attendance.status)])
    rows.sort(key=lambda row: (row[0], row[4], row[5]))
    return _build_pdf("Rekap Presensi Dosen", [("Dosen", current_user.full_name), ("Periode", label)],
        ["Kode", "Mata Kuliah", "Kelas", "Pertemuan", "Tanggal", "NPM", "Mahasiswa", "Status"], rows,
        "statistik-kehadiran.pdf")
